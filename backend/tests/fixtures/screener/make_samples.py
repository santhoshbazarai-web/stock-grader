"""Regenerate the Screener sample workbooks.

Run: ``uv run python tests/fixtures/screener/make_samples.py``.

Layout mirrors Screener's export "Data Sheet" (column A labels, section headers, a
``Report Date`` row per section). Figures are fictional (₹ crore), chosen so FY2024 and the
Mar-2024 quarter can be checked by hand — see test_screener_import.py.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
YEARS = [datetime(y, 3, 31) for y in range(2015, 2025)]
QUARTERS = [
    datetime(2021, 12, 31), datetime(2022, 3, 31), datetime(2022, 6, 30), datetime(2022, 9, 30),
    datetime(2022, 12, 31), datetime(2023, 3, 31), datetime(2023, 6, 30), datetime(2023, 9, 30),
    datetime(2023, 12, 31), datetime(2024, 3, 31),
]  # fmt: skip

# FY2024 values; earlier years are FY2024 / 1.1**(2024 - year), rounded to 2 dp.
PL_2024 = {
    "Sales": 1000, "Raw Material Cost": 400, "Change in Inventory": 10, "Power and Fuel": 50,
    "Other Mfr. Exp": 30, "Employee Cost": 100, "Selling and admin": 60, "Other Expenses": 40,
    "Other Income": 20, "Depreciation": 30, "Interest": 10, "Profit before tax": 310,
    "Tax": 78, "Net profit": 232, "Dividend Amount": 50,
}  # fmt: skip
BS_LIAB_2024 = {"Equity Share Capital": 20, "Reserves": 980, "Borrowings": 150,
                "Other Liabilities": 350, "Total": 1500}  # fmt: skip
BS_ASSET_2024 = {"Net Block": 600, "Capital Work in Progress": 50, "Investments": 200,
                 "Other Assets": 650, "Total": 1500}  # fmt: skip
BS_EXTRA_2024 = {"Receivables": 180, "Inventory": 120, "Cash & Bank": 90}
CF_2024 = {"Cash from Operating Activity": 280, "Cash from Investing Activity": -150,
           "Cash from Financing Activity": -100, "Net Cash Flow": 30}  # fmt: skip
Q_MAR24 = {"Sales": 270, "Expenses": 180, "Other Income": 5, "Depreciation": 8,
           "Interest": 2, "Profit before tax": 85, "Tax": 21, "Net profit": 64,
           "Operating Profit": 90}  # fmt: skip


def scaled(v: float, year: int) -> float:
    return round(v / 1.1 ** (2024 - year), 2)


def block(ws, label_values: dict[str, float], dates: list[datetime], blank: set[tuple[str, int]]):
    for label, v in label_values.items():
        row = [label]
        for d in dates:
            row.append(None if (label, d.year) in blank else scaled(v, d.year))
        ws.append(row)


def build(path: Path, *, shareholding: bool) -> None:
    wb = Workbook()
    wb.active.title = "Profit & Loss"  # other export sheets exist but hold formulas
    for name in ("Quarters", "Balance Sheet", "Cash Flow", "Customization"):
        wb.create_sheet(name)
    ws = wb.create_sheet("Data Sheet")
    ws.append(["COMPANY NAME", "SAMPLE INDUSTRIES LTD"])
    ws.append(["LATEST VERSION", 1.1])
    ws.append(["CURRENT VERSION", 1.1])
    ws.append([])
    ws.append(["META"])
    ws.append(["Number of shares", 200000000])
    ws.append(["Face Value", 1])
    ws.append(["Current Price", 1450])
    ws.append(["Market Capitalization", 29000])
    ws.append([])
    ws.append(["PROFIT & LOSS"])
    ws.append(["Report Date", *YEARS])
    block(ws, PL_2024, YEARS, set())
    ws.append([])
    ws.append(["Quarters"])
    ws.append(["Report Date", *QUARTERS])
    for label, v in Q_MAR24.items():
        ws.append([label, *[round(v * (0.95 ** (len(QUARTERS) - 1 - i)), 2)
                            for i in range(len(QUARTERS))]])  # fmt: skip
    ws.append([])
    ws.append(["BALANCE SHEET"])
    ws.append(["Report Date", *YEARS])
    block(ws, BS_LIAB_2024, YEARS, set())
    block(ws, BS_ASSET_2024, YEARS, set())
    block(ws, BS_EXTRA_2024, YEARS, {("Receivables", 2015)})  # a genuinely missing cell
    ws.append(["No. of Equity Shares", *[200000000] * len(YEARS)])
    ws.append(["New Bonus Shares", *[None] * len(YEARS)])
    ws.append(["Face value", *[1] * len(YEARS)])
    ws.append(["Adjusted Equity Shares in Cr", *[20] * len(YEARS)])
    ws.append([])
    ws.append(["CASH FLOW:"])
    ws.append(["Report Date", *YEARS])
    block(ws, CF_2024, YEARS, set())
    ws.append([])
    ws.append(["PRICE:", *[round(1450 / 1.12 ** (2024 - d.year)) for d in YEARS]])
    ws.append([])
    ws.append(["DERIVED:"])
    ws.append(["Adjusted Equity Shares in Cr", *[20] * len(YEARS)])
    if shareholding:
        ws.append([])
        ws.append(["SHAREHOLDING"])
        dates = [datetime(2023, 6, 30), datetime(2023, 9, 30), datetime(2023, 12, 31),
                 datetime(2024, 3, 31)]  # fmt: skip
        ws.append(["Report Date", *dates])
        ws.append(["Promoters", 55.1, 55.1, 54.9, 54.9])
        ws.append(["FIIs", 17.8, 18.0, 18.4, 18.6])
        ws.append(["DIIs", 12.5, 12.4, 12.6, 12.7])
        ws.append(["Public", 14.6, 14.5, 14.1, 13.8])
        ws.append(["No. of Shareholders", 251000, 248500, 246200, 244900])
    wb.save(path)


if __name__ == "__main__":
    build(HERE / "sample_export.xlsx", shareholding=False)
    build(HERE / "sample_export_with_shareholding.xlsx", shareholding=True)
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(["not", "a", "screener", "export"])
    wb.save(HERE / "not_screener.xlsx")
