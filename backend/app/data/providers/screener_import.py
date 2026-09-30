"""Screener.in Excel export → ``fin_annual`` / ``fin_quarterly`` / ``shareholding``.

Personal use only (SPEC §3.2). The export's "Data Sheet" is a column-A-labelled grid split into
sections by header rows; each section starts with a ``Report Date`` row whose cells are the
period ends for the value columns beneath::

    COMPANY NAME        | Sample Industries Ltd
    PROFIT & LOSS
    Report Date         | 2015-03-31 | ... | 2024-03-31
    Sales               | ...
    Quarters
    Report Date         | 2021-12-31 | ... | 2024-03-31
    BALANCE SHEET / CASH FLOW: / PRICE: / DERIVED: / (optional) SHAREHOLDING

Screener figures are ₹ crore. Labels are looked up via ``app.data.canonical.CANONICAL_FIELDS``
(the only place they are listed); derivations are documented there too.

The standard export has no shareholding block, no announcement dates and no current
liabilities / payables / capex. Those stay NULL and are recorded as data gaps. A
``SHAREHOLDING`` section (added via Screener's export customisation, same layout, labels
Promoters / FIIs / DIIs / Public / No. of Shareholders) is parsed when present.

The export does not say whether figures are consolidated or standalone, so the uploader must
state it (rule 5: consolidated first; standalone is flagged when served).

Exchange results filings (XBRL, ``results_watch``) are the primary source of fundamentals; an
upload never overwrites a period they stored, it only fills fields they lack (e.g. ``sga``) and
periods they don't cover (history before XBRL filing began).
"""

import io
import logging
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import Dataset, Provider
from app.data.canonical import (
    Table,
    fields_for,
    fiscal_year,
    frame_to_rows,
    labels_for,
    pick,
    unavailable_from,
)
from app.data.gaps import GapRecord, GapRecorder
from app.data.providers.base import ProviderUnavailable
from app.data.results_store import merge_upsert
from app.db.base import Base
from app.db.enums import StatementType
from app.db.models import FinAnnual, FinQuarterly, Instrument, Shareholding
from app.db.upsert import upsert

logger = logging.getLogger(__name__)

SOURCE = Provider.SCREENER.value
SHEET = "Data Sheet"
_SECTIONS = {
    "META": "meta",
    "PROFIT & LOSS": "pl",
    "QUARTERS": "quarters",
    "BALANCE SHEET": "bs",
    "CASH FLOW": "cf",
    "PRICE": "price",
    "DERIVED": "derived",
    "SHAREHOLDING": "shp",
    "SHAREHOLDING PATTERN": "shp",
}
_ANNUAL_SECTIONS = ("pl", "bs", "cf", "derived")


class ScreenerFormatError(ValueError):
    """The file is not a Screener export (no Data Sheet, no Report Date rows...)."""


@dataclass
class ScreenerData:
    company_name: str | None
    annual: pd.DataFrame
    quarterly: pd.DataFrame
    shareholding: pd.DataFrame
    warnings: list[str] = field(default_factory=list)


# ───────────────────────── parsing (pure) ─────────────────────────


def _cell_date(value: Any) -> pd.Timestamp | None:
    if value is None or value == "":
        return None
    try:
        ts = pd.Timestamp(value)
    except (ValueError, TypeError):
        return None
    return None if pd.isna(ts) else ts.normalize()


def _cell_num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    return None if pd.isna(f) else f


def _read_sections(rows: list[tuple[Any, ...]]) -> tuple[str | None, dict[str, dict[str, Any]]]:
    """→ (company name, {section: {"dates": [...], "rows": {label: [values...]}}})."""
    company: str | None = None
    sections: dict[str, dict[str, Any]] = {}
    current: str | None = None
    for row in rows:
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip()
        values = list(row[1:])
        header = label.rstrip(":").strip().upper()
        if header == "COMPANY NAME":
            company = str(values[0]).strip() if values and values[0] else None
            continue
        if header in _SECTIONS:
            current = _SECTIONS[header]
            sections.setdefault(current, {"dates": [], "rows": {}})
            continue
        if current is None:
            continue
        sec = sections[current]
        if label.lower() == "report date":
            sec["dates"] = [_cell_date(v) for v in values]
            continue
        sec["rows"].setdefault(label, values)  # first occurrence wins ("Total" appears twice)
    return company, sections


def _by_period(section: dict[str, Any]) -> dict[pd.Timestamp, dict[str, float | None]]:
    out: dict[pd.Timestamp, dict[str, float | None]] = {}
    for i, period in enumerate(section["dates"]):
        if period is None:
            continue
        out[period] = {
            label: _cell_num(values[i]) if i < len(values) else None
            for label, values in section["rows"].items()
        }
    return out


def _sum(*xs: float | None) -> float | None:
    return None if any(x is None for x in xs) else float(sum(x for x in xs if x is not None))


def _derive_annual(rec: dict[str, Any], raw: dict[str, float | None]) -> None:
    """Formulas documented in ``CANONICAL_FIELDS[...].derived["screener"]``."""
    rm, inv_change = raw.get("Raw Material Cost"), raw.get("Change in Inventory")
    if rec.get("cogs") is None and rm is not None and inv_change is not None:
        rec["cogs"] = rm - inv_change
    pbt, interest = rec.get("pbt"), rec.get("interest")
    if rec.get("ebit") is None:
        rec["ebit"] = _sum(pbt, interest)
    if rec.get("ebitda") is None:
        dep, oi = rec.get("depreciation"), rec.get("other_income")
        partial = _sum(pbt, interest, dep)
        rec["ebitda"] = None if partial is None or oi is None else partial - oi
    if rec.get("total_equity") is None:
        rec["total_equity"] = _sum(raw.get("Equity Share Capital"), raw.get("Reserves"))
    shares = rec.get("shares_diluted_cr")
    if shares:
        if rec.get("pat") is not None:
            rec["eps_diluted"] = rec["pat"] / shares
        if rec.get("total_equity") is not None:
            rec["book_value_per_share"] = rec["total_equity"] / shares


def _derive_quarterly(rec: dict[str, Any]) -> None:
    if rec.get("ebit") is None:
        rec["ebit"] = _sum(rec.get("pbt"), rec.get("interest"))


def _canonical(
    periods: dict[pd.Timestamp, dict[str, float | None]],
    table: Table,
    derive: Callable[[dict[str, Any], dict[str, float | None]], None] | None = None,
) -> pd.DataFrame:
    labels = labels_for("screener", table)
    records = {}
    for period_end, raw in sorted(periods.items()):
        rec: dict[str, Any] = {name: pick(raw, lbls) for name, lbls in labels.items()}
        if derive:
            derive(rec, raw)
        if all(v is None for v in rec.values()):
            continue
        if table == "fin_annual":
            rec["fiscal_year"] = fiscal_year(period_end)
        records[period_end] = rec
    cols = fields_for(table) + (["fiscal_year"] if table == "fin_annual" else [])
    df = pd.DataFrame.from_dict(records, orient="index").reindex(columns=cols)
    df.index = pd.DatetimeIndex(df.index, name="period_end")
    return df


def parse_data_sheet(source: bytes | Path | BinaryIO) -> ScreenerData:
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:  # zipfile/openpyxl raise assorted errors for non-xlsx input
        raise ScreenerFormatError(f"not an Excel workbook: {type(exc).__name__}") from exc
    if SHEET not in wb.sheetnames:
        raise ScreenerFormatError(f"no '{SHEET}' sheet (found {wb.sheetnames})")
    rows = [tuple(r) for r in wb[SHEET].iter_rows(values_only=True)]
    wb.close()
    company, sections = _read_sections(rows)
    if "pl" not in sections or not sections["pl"]["dates"]:
        raise ScreenerFormatError("Data Sheet has no PROFIT & LOSS block with Report Date")

    annual_periods: dict[pd.Timestamp, dict[str, float | None]] = {}
    for key in _ANNUAL_SECTIONS:
        if key in sections:
            for period, values in _by_period(sections[key]).items():
                merged = annual_periods.setdefault(period, {})
                for label, v in values.items():
                    if merged.get(label) is None:
                        merged[label] = v

    warnings: list[str] = []
    annual = _canonical(annual_periods, "fin_annual", _derive_annual)
    quarterly = _canonical(
        _by_period(sections["quarters"]) if "quarters" in sections else {},
        "fin_quarterly",
        lambda rec, _raw: _derive_quarterly(rec),
    )
    shareholding = _canonical(
        _by_period(sections["shp"]) if "shp" in sections else {}, "shareholding"
    )
    if "shp" not in sections:
        warnings.append("no SHAREHOLDING block in the export (standard exports omit it)")
    if "quarters" not in sections:
        warnings.append("no Quarters block in the export")
    return ScreenerData(company, annual, quarterly, shareholding, warnings)


# ───────────────────────── persistence ─────────────────────────


@dataclass
class ImportSummary:
    annual_rows: int
    quarterly_rows: int
    shareholding_rows: int
    gaps: list[GapRecord]
    warnings: list[str]


def import_screener(
    session: Session,
    *,
    symbol: str,
    content: bytes | Path | BinaryIO,
    statement_type: StatementType,
    gaps: GapRecorder,
    now: datetime | None = None,
) -> ImportSummary:
    """Parse an export and upsert it for ``symbol`` (which must exist in ``instruments``).
    Records a data gap per canonical field the export cannot supply, plus announcement dates.
    Commits nothing — the caller owns the transaction."""
    instrument_id = session.scalar(select(Instrument.id).where(Instrument.symbol == symbol))
    if instrument_id is None:
        raise ValueError(f"unknown instrument {symbol!r}")
    data = parse_data_sheet(content)
    fetched_at = now or datetime.now(UTC)
    extra = {"statement_type": statement_type}

    annual = data.annual.copy()
    annual["announcement_date"] = None
    quarterly = data.quarterly.copy()
    quarterly["announcement_date"] = None
    counts = {}
    targets: tuple[tuple[type[Base], pd.DataFrame, Table, dict[str, Any] | None], ...] = (
        (FinAnnual, annual, "fin_annual", extra),
        (FinQuarterly, quarterly, "fin_quarterly", extra),
        (Shareholding, data.shareholding, "shareholding", None),
    )
    for model, df, table, cols in targets:
        rows = frame_to_rows(
            df, table, instrument_id=instrument_id, source=SOURCE, fetched_at=fetched_at,
            extra_cols=cols,
        )  # fmt: skip
        if model is FinAnnual or model is FinQuarterly:
            # Periods already stored from exchange filings (XBRL) keep those figures; the
            # upload only fills what the filings lack (see app.data.results_store).
            counts[table] = merge_upsert(session, model, rows, source=SOURCE)
        else:
            counts[table] = upsert(session, model, rows)

    recorded: list[GapRecord] = []
    fin_tables: tuple[tuple[Table, Dataset], ...] = (
        ("fin_annual", Dataset.FIN_ANNUAL),
        ("fin_quarterly", Dataset.FIN_QUARTERLY),
    )
    for table, dataset in fin_tables:
        for name in [*unavailable_from("screener", table), "announcement_date"]:
            recorded.append(GapRecord(dataset, symbol, "not in Screener export", [SOURCE], name))
    if data.shareholding.empty:
        recorded.append(GapRecord(Dataset.SHAREHOLDING, symbol, "not in Screener export", [SOURCE]))
    else:
        for name in unavailable_from("screener", "shareholding"):
            recorded.append(
                GapRecord(Dataset.SHAREHOLDING, symbol, "not in Screener export", [SOURCE], name)
            )
    for gap in recorded:
        gaps.record(gap)
    logger.info("Screener import %s: %s", symbol, counts)
    return ImportSummary(
        counts["fin_annual"], counts["fin_quarterly"], counts["shareholding"], recorded,
        data.warnings,
    )  # fmt: skip


# ───────────────────────── provider (serves uploaded data) ─────────────────────────


class ScreenerProvider:
    """FundamentalsProvider + ShareholdingProvider over previously uploaded Screener data.
    Consolidated first; standalone only when no consolidated upload exists, and flagged
    (rule 5). ``attrs["as_of"]`` is the upload time, so ``staleness_hours`` applies to it."""

    name = Provider.SCREENER

    def __init__(self, session_factory: Callable[[], Session]) -> None:
        self._session_factory = session_factory

    def _frame(
        self, model: type[FinAnnual] | type[FinQuarterly], symbol: str, table: Table
    ) -> pd.DataFrame:
        session = self._session_factory()
        try:
            base = (
                select(model)
                .join(Instrument, Instrument.id == model.instrument_id)
                .where(Instrument.symbol == symbol, model.source == SOURCE)
            )
            chosen: StatementType | None = None
            rows: list[Any] = []
            for basis in (StatementType.CONSOLIDATED, StatementType.STANDALONE):
                rows = list(session.scalars(base.where(model.statement_type == basis)))
                if rows:
                    chosen = basis
                    break
        finally:
            session.close()
        if not rows:
            raise ProviderUnavailable(f"no Screener upload for {symbol}")
        cols = [*fields_for(table), "announcement_date"]
        if table == "fin_annual":
            cols.append("fiscal_year")
        df = pd.DataFrame(
            [{c: getattr(r, c) for c in cols} for r in rows],
            index=pd.DatetimeIndex([pd.Timestamp(r.period_end) for r in rows], name="period_end"),
        ).sort_index()
        df.attrs["as_of"] = max(r.fetched_at for r in rows)
        df.attrs["statement_type"] = chosen.value if chosen else None
        df.attrs["warnings"] = (
            ["standalone figures (no consolidated upload)"]
            if chosen is StatementType.STANDALONE
            else []
        )
        return df

    def annual(self, symbol: str) -> pd.DataFrame:
        return self._frame(FinAnnual, symbol, "fin_annual")

    def quarterly(self, symbol: str) -> pd.DataFrame:
        return self._frame(FinQuarterly, symbol, "fin_quarterly")

    def shareholding(self, symbol: str) -> pd.DataFrame:
        session = self._session_factory()
        try:
            rows = list(
                session.scalars(
                    select(Shareholding)
                    .join(Instrument, Instrument.id == Shareholding.instrument_id)
                    .where(Instrument.symbol == symbol, Shareholding.source == SOURCE)
                )
            )
        finally:
            session.close()
        if not rows:
            raise ProviderUnavailable(f"no Screener shareholding upload for {symbol}")
        cols = [*fields_for("shareholding"), "filing_date"]
        df = pd.DataFrame(
            [{c: getattr(r, c) for c in cols} for r in rows],
            index=pd.DatetimeIndex([pd.Timestamp(r.period_end) for r in rows], name="period_end"),
        ).sort_index()
        df.attrs["as_of"] = max(r.fetched_at for r in rows)
        return df
